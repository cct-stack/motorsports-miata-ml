import os
import numpy as np


class Track:
    """
    Representation of a race track as curvature vs. distance along the line.

    Attributes:
      name       : track name
      s          : distance along the track [m] (fine mesh)
      curvature  : signed curvature [1/m] at each point (+left / -right)
      config     : 'Closed' (loop) or 'Open' (point-to-point)
      length     : total track length [m]
    """

    def __init__(self, name="Skidpad", radius=30.0, n=200):
        # Default: a closed circular skidpad (constant curvature, no straights).
        self.name = name
        self.radius = radius
        self.config = "Closed"
        circumference = 2.0 * np.pi * radius
        self.s = np.linspace(0.0, circumference, n)
        self.curvature = np.ones_like(self.s) / radius
        self.length = circumference

    # ------------------------------------------------------------------ #
    @classmethod
    def _blank(cls, name, s, curvature, config, length):
        obj = cls.__new__(cls)
        obj.name = name
        obj.radius = None
        obj.s = s
        obj.curvature = curvature
        obj.config = config
        obj.length = length
        return obj

    @classmethod
    def from_shape_data(cls, types, lengths, radii, name="Track",
                        config="Closed", mesh_size=2.0):
        """Build a curvature-vs-distance track from OpenTRACK 'shape data'.

        Mirrors OpenTRACK.m: a straight contributes two zero-curvature anchor
        points (its start and end); a corner contributes one anchor at its
        centre with curvature = (+1 left / -1 right) / radius. The anchors are
        then interpolated onto a fine, uniform mesh with a monotone cubic
        (PCHIP) spline, exactly as OpenTRACK does.
        """
        from scipy.interpolate import PchipInterpolator

        types  = list(types)
        lengths = np.asarray(lengths, dtype=float)
        radii  = np.asarray(radii, dtype=float)

        # straight radius -> infinity; sign per corner direction
        sign = np.array([1 if str(t).strip().lower() == "left"
                         else (-1 if str(t).strip().lower() == "right" else 0)
                         for t in types], dtype=float)

        # drop zero-length segments
        keep = lengths > 0
        lengths, radii, sign = lengths[keep], radii[keep], sign[keep]

        # merge consecutive straights so they share single anchors
        m_len, m_sign, m_rad = [], [], []
        for l, s, r in zip(lengths, sign, radii):
            is_straight = (s == 0) or (r == 0)
            if is_straight and m_len and (m_sign[-1] == 0):
                m_len[-1] += l
            else:
                m_len.append(l)
                m_sign.append(0.0 if is_straight else s)
                m_rad.append(np.inf if is_straight else r)

        m_len = np.asarray(m_len)
        X  = np.cumsum(m_len)            # end position of each segment
        XC = X - m_len / 2.0             # centre position of each segment
        L  = float(X[-1])

        xs, rs = [], []
        for i in range(len(m_len)):
            if np.isinf(m_rad[i]):       # straight: anchor start & end at r=0
                xs += [X[i] - m_len[i], X[i]]
                rs += [0.0, 0.0]
            else:                        # corner: anchor centre at r=sign/R
                xs.append(XC[i])
                rs.append(m_sign[i] / m_rad[i])

        # unique, sorted anchors (PCHIP needs strictly increasing x)
        xs = np.asarray(xs)
        rs = np.asarray(rs)
        order = np.argsort(xs, kind="stable")
        xs, rs = xs[order], rs[order]
        uniq = np.concatenate(([True], np.diff(xs) > 1e-9))
        xs, rs = xs[uniq], rs[uniq]

        x_fine = np.arange(0.0, L, mesh_size)
        r_fine = PchipInterpolator(xs, rs, extrapolate=True)(x_fine)
        return cls._blank(name, x_fine, r_fine, config, L)

    @classmethod
    def from_openlap_xlsx(cls, path, mesh_size=2.0):
        """Load an OpenTRACK 'shape data' .xlsx (Info + Shape sheets)."""
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        info = {r[0]: r[1] for r in wb["Info"].iter_rows(values_only=True)
                if r and r[0] is not None}
        name   = str(info.get("Name", os.path.splitext(os.path.basename(path))[0]))
        config = str(info.get("Configuration", "Closed"))

        types, lengths, radii = [], [], []
        for row in wb["Shape"].iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            types.append(row[0])
            lengths.append(float(row[1]))
            radii.append(float(row[2]) if row[2] is not None else 0.0)
        return cls.from_shape_data(types, lengths, radii, name=name,
                                   config=config, mesh_size=mesh_size)

    @classmethod
    def monza(cls, mesh_size=2.0):
        """Autodromo Nazionale Monza, parsed from OpenLAP's own track file so
        our sim and OpenLAP run identical geometry."""
        here = os.path.dirname(os.path.abspath(__file__))
        return cls.from_openlap_xlsx(
            os.path.join(here, "Autodromo Nazionale Monza.xlsx"),
            mesh_size=mesh_size)
